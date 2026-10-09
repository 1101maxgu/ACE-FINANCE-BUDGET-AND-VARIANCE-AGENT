import openpyxl
import pytest
from reportlab.pdfgen import canvas

from variance import explain, gmail, reimburse as R


def form_pdf(path):
    c = canvas.Canvas(str(path))
    c.acroForm.textfield(name="Claimant", value="Ann Lee", x=50, y=700, width=200, height=20)
    c.acroForm.textfield(name="Total Amount", value="$45.50", x=50, y=650, width=200, height=20)
    c.acroForm.textfield(name="Purpose", value="Supplies", x=50, y=600, width=200, height=20)
    c.showPage()
    c.save()


def text_pdf(path, lines):
    c = canvas.Canvas(str(path))
    for i, ln in enumerate(lines):
        c.drawString(50, 750 - 20 * i, ln)
    c.showPage()
    c.save()


def test_fillable_form_profile(tmp_path):
    p = tmp_path / "claim7.pdf"
    form_pdf(p)
    f = R.read_form_fields(p)
    assert f == {"Claimant": "Ann Lee", "Total Amount": "$45.50", "Purpose": "Supplies"}
    prof = R.guess_form_profile(f)
    assert prof["claimant"] == "Claimant" and prof["amount"] == "Total Amount" and prof["description"] == "Purpose"
    R.save_form_profile("Alice form", prof, tmp_path / "fp.yaml")
    assert R.load_form_profiles(tmp_path / "fp.yaml")["Alice form"] == prof
    c = R.read_claim_form(p, prof)
    assert (c["claim_id"], c["claimant"], c["amount_cents"], c["confidence"]) == ("claim7", "Ann Lee", 4550, 1.0)


def test_parse_receipt_prefers_total_line_and_keeps_its_confidence():
    r = R.parse_receipt([("ACME Store", 0.99), ("Subtotal 40.00", 0.95), ("TOTAL $1,045.50", 0.6), ("Date 2026-01-05", 0.9)])
    assert (r["amount_cents"], r["confidence"], r["date"], r["vendor"]) == (104550, 0.6, "2026-01-05", "ACME Store")
    assert R.parse_receipt([("x", 1), ("12.00", 0.9), ("30.25", 0.8)])["amount_cents"] == 3025   # no total line: largest
    assert R.parse_receipt([("Paid 03/04/2026 7.00", 0.9)])["date"] is None                     # ambiguous: not guessed
    assert R.parse_receipt([("no money here", 0.9)])["amount_cents"] is None


def test_low_confidence_does_not_count_until_confirmed():
    low = {"receipt_id": "r", "amount_cents": 1000, "confidence": 0.5, "confirmed": False}
    assert not R.counts(low, 0.8) and R.counts({**low, "confidence": 0.8}, 0.8)
    ok = R.confirm(low, "Ann", amount_cents=1200)
    assert R.counts(ok, 0.8) and ok["amount_cents"] == 1200 and ok["confirmed_by"] == "Ann" and low["confirmed"] is False
    with pytest.raises(ValueError):
        R.confirm(low, " ")


def rc(rid, claim, cents, conf=0.95, date="2026-01-05", vendor="V", confirmed=False):
    return {"receipt_id": rid, "claim_id": claim, "amount_cents": cents, "confidence": conf, "date": date,
            "vendor": vendor, "confirmed": confirmed}


CLAIMS = [
    {"claim_id": "C1", "amount_cents": 4550, "items": [4000, 550]},
    {"claim_id": "C2", "amount_cents": 3000},
    {"claim_id": "C3", "amount_cents": 2000},
    {"claim_id": "C4", "amount_cents": 1000},
    {"claim_id": "C5", "amount_cents": 5000, "items": [2000, 2000]},
    {"claim_id": "C6", "amount_cents": 700},
]
RECEIPTS = [rc("R1", "C1", 4550), rc("R3", "C3", 2000, conf=0.5), rc("R4a", "C4", 1000, vendor="Shop"),
            rc("R4b", "C4", 1000, vendor="shop"), rc("R5", "C5", 5000, date="2026-01-06"),
            rc("R6", "C6", 500, date="2026-01-07"), rc("R9", None, 300, date="2026-01-08")]


def test_reconcile_all_statuses():
    out = R.reconcile(CLAIMS, RECEIPTS, 0.8)
    by = {r["claim_id"]: r for r in out["claims"]}
    assert by["C1"]["status"] == "matched" and by["C1"]["receipts_cents"] == 4550
    assert by["C2"]["status"] == "missing_receipt"
    assert by["C3"]["status"] == "amount_mismatch" and by["C3"]["receipts_cents"] == 0
    assert by["C3"]["needs_confirmation"] == ["R3"]
    assert by["C4"]["status"] == "duplicate_receipt" and by["C4"]["duplicate_ids"] == ["R4b"] and by["C4"]["receipts_cents"] == 1000
    assert by["C5"]["status"] == "total_mismatch" and by["C5"]["statuses"] == ["total_mismatch"]
    assert by["C6"]["status"] == "amount_mismatch" and by["C6"]["receipts_cents"] == 500
    assert out["unclaimed"] == [{"receipt_id": "R9", "status": "unclaimed_receipt", "amount_cents": 300}]
    # a person confirms the low-confidence receipt -> now it counts
    fixed = [R.confirm(r, "Ann") if r["receipt_id"] == "R3" else r for r in RECEIPTS]
    assert {r["claim_id"]: r["status"] for r in R.reconcile(CLAIMS, fixed, 0.8)["claims"]}["C3"] == "matched"


def test_duplicate_by_file_hash():
    a, b = {**rc("A", "C1", 100), "sha256": "x", "date": None}, {**rc("B", "C1", 100), "sha256": "x", "date": None}
    assert R.reconcile([{"claim_id": "C1", "amount_cents": 100}], [a, b])["claims"][0]["duplicate_ids"] == ["B"]


def test_text_pdf_receipt_and_missing_tesseract(tmp_path, monkeypatch):
    p = tmp_path / "r.pdf"
    text_pdf(p, ["Corner Shop", "Subtotal 40.00", "Total 45.50"])
    r = R.read_receipt(p, claim_id="C1")
    assert (r["amount_cents"], r["source"], r["confidence"], r["claim_id"], r["confirmed"]) == (4550, "text", 1.0, "C1", False)
    # a scan (no text layer) needs OCR; pretend Tesseract is not installed
    import pytesseract
    monkeypatch.setattr(pytesseract, "get_tesseract_version",
                        lambda: (_ for _ in ()).throw(pytesseract.TesseractNotFoundError()))
    blank = tmp_path / "scan.pdf"
    text_pdf(blank, [])
    with pytest.raises(R.OcrUnavailable, match="Tesseract OCR"):
        R.read_receipt(blank)


def test_claim_store_states_and_committed(tmp_path):
    st = {"shared_folder": str(tmp_path)}
    R.upsert_claim(st, {"claim_id": "C1", "amount_cents": 4550, "line_id": "L1"})
    R.upsert_claim(st, {"claim_id": "C2", "amount_cents": 1000, "line_id": "L1"})
    assert R.load_claims(st)["C1"]["state"] == "received"
    assert R.committed_by_line(st) == {}
    with pytest.raises(ValueError, match="person"):
        R.set_state(st, "C1", "approved", "system")
    with pytest.raises(ValueError):
        R.set_state(st, "C1", "approved", "")
    R.set_state(st, "C1", "approved", "Dana", date="2026-02-01")
    R.upsert_claim(st, {"claim_id": "C1", "amount_cents": 4550, "line_id": "L1"})        # refresh keeps state
    assert R.load_claims(st)["C1"]["state"] == "approved"
    assert R.committed_by_line(st) == {"L1": 4550}
    R.set_state(st, "C1", "paid", "Dana")
    assert R.committed_by_line(st) == {}


def test_suggested_state_is_never_approved():
    out = R.reconcile(CLAIMS, RECEIPTS, 0.8)["claims"]
    assert {R.suggested_state(r) for r in out} <= {"ready_to_approve", "needs_info"}
    assert [R.suggested_state(r) for r in out][0] == "ready_to_approve"


def test_weekly_summary_text_is_guard_clean_and_excel(tmp_path):
    st = {"shared_folder": str(tmp_path)}
    out = R.reconcile(CLAIMS, RECEIPTS, 0.8)
    for c in CLAIMS:
        R.upsert_claim(st, c)
    R.set_state(st, "C1", "approved", "Dana")
    s = R.weekly_summary(st, out["claims"], "Jan 5-11")
    # claimed: 4550+3000+2000+1000+5000+700 = 16250; ready: only C1 (matched) = 4550; committed 4550
    assert (s["claims_count"], s["claimed_cents"], s["ready_cents"], s["committed_cents"]) == (6, 16250, 4550, 4550)
    assert s["by_state"]["approved"] == 1 and s["by_status"]["amount_mismatch"] == 2
    txt = R.claims_text(s)
    assert "$162.50 claimed" in txt and "$45.50 are ready" in txt
    assert explain.number_guard(txt, s) == []
    assert explain.number_guard(txt + " Also $999.00.", s) == ["999.00"]
    mail = R.exception_email({"claim_id": "C2", "claimant": "Ann", "email": "a@x.org"}, out["claims"][1])
    assert mail["to"] == "a@x.org" and "attach the receipt" in mail["subject"]
    c = gmail.InMemoryGmailClient()
    gmail.make_draft(c, mail["to"], mail["subject"], mail["body"])
    assert len(c.drafts) == 1 and c.sent == []
    x = R.write_claims_excel(CLAIMS, out["claims"], out["unclaimed"], tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(x)
    assert wb["Claims"].max_row == 7 and wb["Exceptions"].max_row == 7    # header + 6 claims; header + 5 exception claims + 1 unclaimed
