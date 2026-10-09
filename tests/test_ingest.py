import openpyxl
import pytest

from variance import ingest
from variance.ingest import IngestError, load_actuals, load_budget, to_cents


def write(tmp_path, name, text):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def test_to_cents():
    assert to_cents("$1,234.56") == 123456
    assert to_cents("(12.50)") == -1250
    assert to_cents("-0.005") == -1  # half-up away from zero
    assert to_cents(12.5) == 1250
    assert to_cents("") is None
    with pytest.raises(ValueError):
        to_cents("abc")


def test_csv_notes_above_header_and_total_row_skipped(tmp_path):
    p = write(tmp_path, "b.csv",
              "Budget FY26,,\nprepared by Sam,,\nItem,Group,Amount\nVenue,Events,\"$1,000.50\"\nFood,Events,200\nTotal,,1200.50\n")
    assert ingest.guess_header_row(p) == 2
    prof = {"header_row": 2, "columns": {"line_name": "Item", "category": "Group", "budget": "Amount"}}
    df = load_budget(p, prof)
    assert list(df.line_id) == ["L001", "L002"]
    assert list(df.budget_cents) == [100050, 20000]
    assert list(df.source_row) == [4, 5]          # 1-based row in the file
    assert df.attrs["skipped_rows"] == [(6, "Total")]
    assert list(ingest.preview(p, header_row=2, n=1).columns) == ["Item", "Group", "Amount"]


def test_budget_missing_column_names_it(tmp_path):
    p = write(tmp_path, "b.csv", "Item,Amount\nVenue,5\n")
    with pytest.raises(IngestError, match="Budgeted"):
        load_budget(p, {"columns": {"line_name": "Item", "budget": "Budgeted"}})


def test_budget_unmapped_required_field(tmp_path):
    p = write(tmp_path, "b.csv", "Item,Amount\nVenue,5\n")
    with pytest.raises(IngestError, match="budget amount"):
        load_budget(p, {"columns": {"line_name": "Item"}})


def test_excel_sections_and_merged_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Line", "Amount", "Group"])
    ws.append(["Event A"])
    ws.append(["Venue", 100.5, "Fixed"])
    ws.append(["Food", 50])
    ws.append(["Event B"])
    ws.append(["Venue", 70])
    ws.merge_cells("C3:C4")          # merged group cell applies to both rows
    wb.create_sheet("Other")
    p = tmp_path / "b.xlsx"
    wb.save(p)
    assert ingest.list_sheets(p) == ["Plan", "Other"]
    cols = {"line_name": "Line", "budget": "Amount"}
    df = load_budget(p, {"sheet": "Plan", "layout": "sections", "columns": cols})
    assert list(zip(df.category, df.line_name, df.budget_cents)) == [
        ("Event A", "Venue", 10050), ("Event A", "Food", 5000), ("Event B", "Venue", 7000)]
    df2 = load_budget(p, {"sheet": "Plan", "columns": {**cols, "category": "Group"}})
    assert df2.category[1:3].tolist() == ["Fixed", "Fixed"]   # rows 3-4 (row 2 is the section title line)
    with pytest.raises(IngestError, match="Nope"):
        load_budget(p, {"sheet": "Nope", "columns": cols})


def test_monthly_wide_sums_months(tmp_path):
    p = write(tmp_path, "b.csv", "Line,Jan,Feb,Mar\nRent,100,100,100\nAds,10.50,0,\n")
    prof = {"layout": "monthly_wide", "month_columns": ["Jan", "Feb", "Mar"], "columns": {"line_name": "Line"}}
    assert list(load_budget(p, prof).budget_cents) == [30000, 1050]


def test_duplicate_line_ids_rejected(tmp_path):
    p = write(tmp_path, "b.csv", "Id,Item,Amt\nA,x,1\nA,y,2\n")
    with pytest.raises(IngestError, match="'A'"):
        load_budget(p, {"columns": {"line_id": "Id", "line_name": "Item", "budget": "Amt"}})


ACT = "Date,Cat,Amt,Memo\n2026-01-05,Food,-12.30,lunch\n2026-01-06,Food,5.00,refund\n"


def test_actuals_sign_conventions(tmp_path):
    p = write(tmp_path, "a.csv", ACT)
    cols = {"date": "Date", "category": "Cat", "amount": "Amt", "description": "Memo"}
    neg = load_actuals(p, {"sign": "negative_expenses", "columns": cols})
    assert list(neg.amount_cents) == [1230, -500]       # expense positive, refund negative
    pos = load_actuals(p, {"sign": "positive_expenses", "columns": cols})
    assert list(pos.amount_cents) == [-1230, 500]
    assert list(neg.txn_id) == ["T0001", "T0002"] and list(neg.description) == ["lunch", "refund"]


def test_actuals_debit_credit(tmp_path):
    p = write(tmp_path, "a.csv", "Date,Cat,Out,In\n2026-01-05,Food,10.00,\n2026-01-06,Food,,2.50\n")
    df = load_actuals(p, {"sign": "debit_credit",
                          "columns": {"date": "Date", "category": "Cat", "debit": "Out", "credit": "In"}})
    assert list(df.amount_cents) == [1000, -250]


def test_ambiguous_dates_flagged_not_guessed(tmp_path):
    cols = {"date": "Date", "category": "Cat", "amount": "Amt"}
    p = write(tmp_path, "a.csv", "Date,Cat,Amt\n03/04/2026,Food,1\n")
    with pytest.raises(IngestError, match="day-first or month-first"):
        load_actuals(p, {"columns": cols})
    assert load_actuals(p, {"columns": cols, "date_order": "dmy"}).date[0].month == 4
    assert load_actuals(p, {"columns": cols, "date_order": "mdy"}).date[0].month == 3
    # another row proves the file is day-first, so the ambiguous row is not a guess
    p2 = write(tmp_path, "a2.csv", "Date,Cat,Amt\n03/04/2026,Food,1\n25/04/2026,Food,1\n")
    assert load_actuals(p2, {"columns": cols}).date[0].month == 4


def test_actuals_errors_name_row_and_column(tmp_path):
    p = write(tmp_path, "a.csv", "Date,Cat,Amt\n2026-01-05,Food,abc\n")
    with pytest.raises(IngestError, match="row 2.*amount"):
        load_actuals(p, {"columns": {"date": "Date", "category": "Cat", "amount": "Amt"}})
    with pytest.raises(IngestError, match="'When'"):
        load_actuals(p, {"columns": {"date": "When", "category": "Cat", "amount": "Amt"}})


def test_file_like_input(tmp_path):
    import io
    f = io.BytesIO(ACT.encode())
    f.name = "x.csv"
    df = load_actuals(f, {"sign": "positive_expenses", "columns": {"date": "Date", "category": "Cat", "amount": "Amt"}})
    assert len(df) == 2
    assert len(load_actuals(f, {"sign": "positive_expenses",
                                "columns": {"date": "Date", "category": "Cat", "amount": "Amt"}})) == 2  # re-readable
