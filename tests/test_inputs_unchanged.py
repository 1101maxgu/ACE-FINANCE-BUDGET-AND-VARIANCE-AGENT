"""Non-negotiable: source files are byte-identical after a full run (and reports are written elsewhere)."""
import datetime as dt
import hashlib

import openpyxl
from openpyxl import load_workbook

from variance import explain, facts as F, ingest, reports
from variance.calc import compute


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def test_full_run_leaves_inputs_untouched(tmp_path):
    src = tmp_path / "inputs"
    src.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Notes: draft"])
    ws.append(["Line", "Cat", "Budget"])
    ws.append(["Food", "Food", 800])
    ws.append(["Venue", "Venue", 500])
    wb.save(src / "budget.xlsx")
    (src / "actuals.csv").write_text("Date,Category,Amount\n2026-01-05,Food,-1100\n2026-01-06,Parking,-200\n")
    before = {p.name: sha(p) for p in src.iterdir()}

    b = ingest.load_budget(src / "budget.xlsx", {"header_row": 1, "columns": {
        "line_name": "Line", "category": "Cat", "budget": "Budget"}})
    a = ingest.load_actuals(src / "actuals.csv", {"sign": "negative_expenses", "columns": {
        "date": "Date", "category": "Category", "amount": "Amount"}})
    ingest.list_sheets(src / "budget.xlsx")
    ingest.preview(src / "actuals.csv")
    r = compute(b, a, {}, {"flag_pct": 10, "flag_min_cents": 15000}, dt.date(2026, 1, 31))
    facts = F.build_facts(r, {})
    out = tmp_path / "out"
    out.mkdir()
    reports.write_excel(r, facts, out / "d.xlsx")
    reports.write_pdf(explain.explain_template(facts), facts, out / "s.pdf")

    assert {p.name: sha(p) for p in src.iterdir()} == before
    assert (out / "s.pdf").read_bytes()[:4] == b"%PDF"
    wb = load_workbook(out / "d.xlsx")
    assert wb.sheetnames == ["Summary", "Lines", "Transactions"]
    lines = wb["Lines"]
    assert lines.freeze_panes == "A2"
    rows = {row[1].value: row for row in lines.iter_rows(min_row=2)}
    assert rows["L001"][0].value == 3 and rows["L001"][7].value == 1100.0     # source row 3, actual $1,100.00
    assert rows["UNBUDGETED:Parking"][0].value is None
    assert rows["L001"][10].fill.start_color.rgb.endswith("FFC7CE")           # unfavorable colored
    assert wb["Transactions"].max_row == 3
