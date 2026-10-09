"""Excel detail (openpyxl) and one-page PDF (reportlab)."""
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

FOOTER = "Generated, not audited."
_FILL = {"favorable": "C6EFCE", "unfavorable": "FFC7CE"}
_MONEY, _PCT = '#,##0.00;[Red]-#,##0.00', '0.00"%"'


def _usd(cents):
    c = abs(int(cents))
    return f"{'-' if cents < 0 else ''}${c // 100:,}.{c % 100:02d}"


def _sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, len(h) + 2))


def write_excel(result, facts, out_path):
    """Summary + line-by-line sheet (with source rows, status colors, Notes column) + transactions."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    t = facts["totals"]
    ws.append(["Budget & Variance report" + ("  -  UNRECONCILED" if facts.get("unreconciled") else "")])
    ws["A1"].font = Font(bold=True, size=14, color="C00000" if facts.get("unreconciled") else "000000")
    for k, v in [("Period", facts["period"].get("label") or f"{facts['period'].get('start')} to {facts['period'].get('end')}"),
                 ("As of", facts["as_of_date"]),
                 ("Flag rule", f"|variance| >= {_usd(facts['thresholds']['flag_min_cents'])} AND >= {facts['thresholds']['flag_pct']}%"),
                 ("Expense budget", t["expense_budget_cents"] / 100), ("Expense actual", t["expense_actual_cents"] / 100),
                 ("Revenue budget", t["revenue_budget_cents"] / 100), ("Revenue actual", t["revenue_actual_cents"] / 100),
                 ("Net variance (positive = favorable)", t["net_variance_cents"] / 100),
                 ("Flagged lines", t["flagged_count"]),
                 ("Unbudgeted spend", t["unbudgeted_cents"] / 100),
                 ("Committed (approved, unpaid claims)", t["committed_cents"] / 100),
                 ("Rows outside period (excluded)", t["excluded_row_count"])] + \
                [("Quality: " + f["severity"], f["message"]) for f in facts.get("quality_findings", [])] + \
                [("", ""), ("", FOOTER)]:
        ws.append([k, v])
        if isinstance(v, float):
            ws.cell(ws.max_row, 2).number_format = _MONEY
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 28

    L = result.lines
    ws = wb.create_sheet("Lines")
    head = ["Source row", "Line ID", "Category", "Line", "Type", "Owner", "Budget", "Actual", "Variance",
            "Variance %", "Status", "Flagged", "Txns", "Notes", "Committed", "Expected to date", "Pace variance",
            "Projected", "Projection confidence", "Chip"]
    d100 = lambda v: None if pd.isna(v) else int(v) / 100
    _sheet(ws, head, [[None if pd.isna(r.budget_source_row) else int(r.budget_source_row), r.line_id,
                       r.category, r.line_name, r.type, r.owner, r.budget_cents / 100, r.actual_cents / 100,
                       r.variance_cents / 100, None if r.variance_pct != r.variance_pct else r.variance_pct,
                       r.status, "YES" if r.flagged else "", r.txn_count, "", r.committed_cents / 100,
                       d100(r.expected_cents), d100(r.pace_variance_cents), d100(r.projection_cents),
                       r.projection_confidence, r.chip] for r in L.itertuples()],
           {"Line": 28, "Category": 20, "Notes": 40})
    for row in range(2, ws.max_row + 1):
        for col in (7, 8, 9, 15, 16, 17, 18):
            ws.cell(row, col).number_format = _MONEY
        ws.cell(row, 10).number_format = _PCT
        fill = _FILL.get(ws.cell(row, 11).value)
        if fill:
            ws.cell(row, 11).fill = PatternFill("solid", start_color=fill)

    ws = wb.create_sheet("Transactions")
    red = facts.get("redacted")
    _sheet(ws, ["Source row", "Txn ID", "Date", "Amount (out +)", "Line ID", "Category", "Vendor", "Description"],
           [[None if pd.isna(x.source_row) else int(x.source_row), x.txn_id, x.date.date(), x.amount_cents / 100, x.line_id, x.category,
             "" if red else x.vendor, "" if red else x.description] for x in result.txns.itertuples()],
           {"Description": 40, "Date": 12})
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 3).number_format = "yyyy-mm-dd"
        ws.cell(row, 4).number_format = _MONEY
    wb.save(out_path)
    return str(Path(out_path))


def _text(item):
    if isinstance(item, dict):
        return f"<b>{escape(str(item.get('title', '')))}</b>. {escape(str(item.get('detail', '')))}"
    return escape(str(item))


def write_pdf(explanation_dict, facts, out_path):
    """One page: status headline, up to 5 issues, decisions needed, quality caveat, footer."""
    ss = getSampleStyleSheet()
    body = ss["BodyText"]
    body.fontSize, body.leading = 9.5, 12
    story = [Paragraph("Budget &amp; Variance summary", ss["Title"])]
    p = facts["period"]
    story.append(Paragraph(escape(f"{p.get('label') or (str(p.get('start')) + ' to ' + str(p.get('end')))} "
                                  f"(as of {facts['as_of_date']})"), body))
    story += [Spacer(1, 8), Paragraph(f"<b>{escape(str(explanation_dict['headline']))}</b>", ss["Heading3"])]
    story.append(Paragraph("Top issues", ss["Heading4"]))
    for it in explanation_dict["top_issues"][:5]:
        story.append(Paragraph("&bull; " + _text(it), body))
        if isinstance(it, dict):
            c = it.get("cause") or {}
            story.append(Paragraph(f"&nbsp;&nbsp;<i>Cause ({escape(str(c.get('label', '?')))}): "
                                   f"{escape(str(c.get('text', '')))}</i>", body))
    story.append(Paragraph("Decisions needed", ss["Heading4"]))
    for d in explanation_dict["decisions_needed"][:6] or ["None."]:
        story.append(Paragraph("&bull; " + escape(str(d)), body))
    q = facts.get("quality_findings") or []
    if facts.get("unreconciled"):
        story.insert(1, Paragraph('<font color="#C00000"><b>UNRECONCILED: the actuals do not match the stated total. '
                                  'Do not rely on these figures yet.</b></font>', body))
    story += [Spacer(1, 8), Paragraph(
        "<i>Data quality: " + (f"{len(q)} finding(s) noted; see detail workbook." if q else "no problems found.") +
        f" {facts['totals']['excluded_row_count']} row(s) fell outside the period and were left out.</i>", body)]

    def foot(canvas, doc):
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.7 * inch, 0.5 * inch, FOOTER + " Numbers come from code; explanations are not an audit.")

    SimpleDocTemplate(str(out_path), pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                      topMargin=0.7 * inch, bottomMargin=0.8 * inch).build(story, onFirstPage=foot)
    return str(Path(out_path))
