"""Read budget / actuals files (CSV or Excel) into normalized frames.

Inputs are only ever read as bytes (never opened for writing).
Money -> integer cents. Actuals sign: normalized amount_cents is POSITIVE for money
out (cost) and NEGATIVE for money in (refund / revenue).
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from dateutil import parser as dateparser

BUDGET_COLS = ["line_id", "category", "line_name", "type", "budget_cents", "owner"]
ACTUAL_COLS = ["txn_id", "date", "amount_cents", "category", "description", "vendor"]
_TOTAL_ROW = re.compile(r"^\s*(grand\s+total|sub\s*-?\s*total|total)\s*:?\s*$", re.I)


class IngestError(Exception):
    """Plain-language problem with an input file (names the column / row)."""


# ---------- reading ----------

def _bytes(file):
    if hasattr(file, "getvalue"):
        return file.getvalue(), getattr(file, "name", "")
    if hasattr(file, "read"):
        if hasattr(file, "seek"):
            file.seek(0)
        return file.read(), getattr(file, "name", "")
    p = Path(file)
    with open(p, "rb") as f:  # read-only
        return f.read(), p.name


def _is_xlsx(data, name):
    n = name.lower()
    if n.endswith(".xls"):
        raise IngestError("Old .xls files are not supported. Please re-save the file as .xlsx or .csv.")
    return n.endswith((".xlsx", ".xlsm")) or (not n.endswith((".csv", ".txt")) and data[:2] == b"PK")


def list_sheets(file) -> list:
    """Sheet names of an Excel file; ['(csv)'] for CSV."""
    data, name = _bytes(file)
    if not _is_xlsx(data, name):
        return ["(csv)"]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _read_rows(file, sheet=None):
    data, name = _bytes(file)
    if _is_xlsx(data, name):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        if sheet and sheet != "(csv)":
            if sheet not in wb.sheetnames:
                raise IngestError(f"Sheet '{sheet}' not found. Sheets in this file: {', '.join(wb.sheetnames)}")
            ws = wb[sheet]
        else:
            ws = wb.active
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:  # fill merged cells so every row carries the value
            v = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[r - 1][c - 1] = v
        wb.close()
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        grid = list(csv.reader(io.StringIO(text), dialect))
    width = max((len(r) for r in grid), default=0)
    return [[("" if v is None else (v.strip() if isinstance(v, str) else v)) for v in r] + [""] * (width - len(r))
            for r in grid]


def _blank(v):
    return v == "" or (isinstance(v, float) and v != v)


def _header_text(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _table(file, sheet=None, header_row=0):
    """Raw cells with the chosen header row applied. Index = 1-based file/sheet row number."""
    rows = _read_rows(file, sheet)
    if header_row >= len(rows):
        raise IngestError(f"Header row {header_row + 1} is past the end of the file ({len(rows)} rows).")
    heads, seen = [], {}
    for i, v in enumerate(rows[header_row]):
        h = _header_text(v) or f"Unnamed: {i + 1}"
        seen[h] = seen.get(h, 0) + 1
        heads.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    body = [(i + 1, r) for i, r in enumerate(rows) if i > header_row and not all(_blank(v) for v in r)]
    df = pd.DataFrame([r for _, r in body], columns=heads, dtype=object)
    df.index = [i for i, _ in body]
    return df


def preview(file, sheet=None, header_row=0, n=10) -> pd.DataFrame:
    """First n data rows (as text) under the chosen header row, for the mapping UI."""
    return _table(file, sheet, header_row).head(n).astype(str)


def guess_header_row(file, sheet=None) -> int:
    """0-based index of the likely header row (skips notes above the table)."""
    rows = _read_rows(file, sheet)[:30]
    counts = [sum(not _blank(v) for v in r) for r in rows]
    if not counts or max(counts) == 0:
        return 0
    need = max(2, 0.6 * max(counts))
    for i, r in enumerate(rows):
        if counts[i] >= need and all(isinstance(v, str) for v in r if not _blank(v)):
            return i
    return 0


# ---------- value parsing ----------

def to_cents(v):
    """'$1,234.56' / '(12.50)' / 12.5 -> integer cents. Blank -> None. Bad text -> ValueError."""
    if v is None or _blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError(v)
    if isinstance(v, str):
        s = v.strip()
        neg = (s.startswith("(") and s.endswith(")")) or s.startswith("-") or s.endswith("-")
        s = re.sub(r"[$£€,\s()\-]", "", s)
        try:
            d = Decimal(s)
        except InvalidOperation:
            raise ValueError(v)
        d = -d if neg else d
    else:
        d = Decimal(str(v))
    return int((d * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


_NUM_DATE = re.compile(r"(\d{1,4})[-/.](\d{1,2})[-/.](\d{1,4})")


def _parse_dates(vals, order, who):
    """Parse a list of cell values to dates. order: 'dmy' | 'mdy' | None (infer, else error)."""
    parsed, pending, evidence = [None] * len(vals), [], set()
    for i, v in enumerate(vals):
        if isinstance(v, datetime):
            parsed[i] = v.date()
        elif isinstance(v, date):
            parsed[i] = v
        elif isinstance(v, str):
            s = v.strip()
            m = _NUM_DATE.fullmatch(s[:10]) if re.match(r"\d{4}-\d{2}-\d{2}", s) else _NUM_DATE.fullmatch(s)
            if not m:
                try:
                    parsed[i] = dateparser.parse(s).date()
                except (ValueError, OverflowError):
                    raise IngestError(f"{who} row {i + 1}: I can't read the date '{v}'.")
                continue
            a, b, c = m.groups()
            try:
                if len(a) == 4:
                    parsed[i] = date(int(a), int(b), int(c))
                    continue
                y = int(c) + (2000 if len(c) <= 2 else 0)
                x, z = int(a), int(b)
                if x > 12:
                    evidence.add("dmy"); parsed[i] = date(y, z, x)
                elif z > 12:
                    evidence.add("mdy"); parsed[i] = date(y, x, z)
                elif x == z:
                    parsed[i] = date(y, x, z)
                else:
                    pending.append((i, x, z, y))
            except ValueError:
                raise IngestError(f"{who} row {i + 1}: '{v}' is not a real date.")
        else:
            raise IngestError(f"{who} row {i + 1}: I can't read the date '{v}'.")
    if pending:
        if order is None and len(evidence) == 1:
            order = next(iter(evidence))
        if len(evidence) > 1 and order is None:
            raise IngestError(f"{who}: dates mix day-first and month-first formats. Please fix the file.")
        if order is None:
            raise IngestError(
                f"{who}: dates like '{pending[0][1]}/{pending[0][2]}/{pending[0][3]}' could be day-first or "
                "month-first and I won't guess. Tell me which one this file uses (date order setting).")
        for i, x, z, y in pending:
            try:
                parsed[i] = date(y, z, x) if order == "dmy" else date(y, x, z)
            except ValueError:
                raise IngestError(f"{who} row {i + 1}: date is not valid as {order}.")
    return parsed


def _sub(profile, kind):
    return profile[kind] if kind in profile and isinstance(profile[kind], dict) else profile


def _pick(df, cols, field, what, who, required=False):
    name = cols.get(field)
    if not name:
        if required:
            raise IngestError(f"{who}: no column chosen for {what}. Please pick one in the mapping step.")
        return None
    if name not in df.columns:
        raise IngestError(f"{who}: I can't find the column '{name}' (needed for {what}). "
                          f"Columns I can see: {', '.join(map(str, df.columns))}")
    return name


def _money(v, who, row, what):
    try:
        return to_cents(v)
    except ValueError:
        raise IngestError(f"{who} row {row}: '{v}' in {what} is not a number.")


def _norm_type(v, default):
    s = str(v).strip().lower()
    if not s:
        return default
    return "revenue" if s.startswith(("rev", "inc")) else "expense"


# ---------- budget ----------

def load_budget(file, profile: dict) -> pd.DataFrame:
    """Normalized budget: line_id, category, line_name, type, budget_cents, owner (+ source_row).

    profile (or profile['budget']): sheet, header_row, layout ('total'|'monthly_wide'|'sections'),
    columns{line_id,category,line_name,type,budget,owner}, month_columns[], type_default.
    Rows that look like totals are skipped and listed in df.attrs['skipped_rows'].
    """
    p = _sub(profile, "budget")
    who = "Budget file"
    df = _table(file, p.get("sheet"), p.get("header_row", 0))
    cols = p.get("columns", {})
    layout = p.get("layout", "total")
    c_name = _pick(df, cols, "line_name", "the line name", who, required=True)
    c_id, c_cat = _pick(df, cols, "line_id", "line id", who), _pick(df, cols, "category", "category", who)
    c_type, c_own = _pick(df, cols, "type", "type", who), _pick(df, cols, "owner", "owner", who)
    if layout == "monthly_wide":
        months = p.get("month_columns") or []
        if not months:
            raise IngestError(f"{who}: monthly layout needs the month columns chosen.")
        amt_cols = [_pick(df, {"m": m}, "m", "a month amount", who) for m in months]
    else:
        amt_cols = [_pick(df, cols, "budget", "the budget amount", who, required=True)]
    tdef = p.get("type_default", "expense")
    recs, skipped, section = [], [], ""
    for row, r in df.iterrows():
        name = str(r[c_name]).strip()
        amts = [r[c] for c in amt_cols]
        if layout == "sections" and name and all(_blank(a) for a in amts):
            section = name
            continue
        if not name and all(_blank(a) for a in amts):
            continue
        if _TOTAL_ROW.match(name):
            skipped.append((row, name))
            continue
        parts = [_money(a, who, row, "the budget amount") or 0 for a in amts]
        cents = sum(parts)
        cat = str(r[c_cat]).strip() if c_cat else ""
        cat = cat or section or name
        recs.append({
            "line_id": str(r[c_id]).strip() if c_id and str(r[c_id]).strip() else f"L{len(recs) + 1:03d}",
            "category": cat, "line_name": name,
            "type": _norm_type(r[c_type], tdef) if c_type else tdef,
            "budget_cents": cents, "owner": str(r[c_own]).strip() if c_own else "", "source_row": row,
            "phasing": list(zip(months, parts)) if layout == "monthly_wide" else None})
    out = pd.DataFrame(recs, columns=BUDGET_COLS + ["source_row", "phasing"])
    if layout != "monthly_wide":
        out = out.drop(columns="phasing")
    dup = out[out.line_id.duplicated()].line_id.tolist()
    if dup:
        raise IngestError(f"{who}: line id '{dup[0]}' appears more than once. Line ids must be unique.")
    out["budget_cents"] = out["budget_cents"].astype("int64")
    out.attrs["skipped_rows"] = skipped
    return out


# ---------- actuals ----------

def load_actuals(file, profile: dict) -> pd.DataFrame:
    """Normalized actuals: txn_id, date, amount_cents, category, description, vendor (+ source_row).

    profile (or profile['actuals']): sheet, header_row, sign ('negative_expenses'|'positive_expenses'|
    'debit_credit'), date_order (None|'dmy'|'mdy'), columns{txn_id,date,amount,debit,credit,category,
    description,vendor}. amount_cents: positive = money out, negative = money in.
    """
    p = _sub(profile, "actuals")
    who = "Actuals file"
    df = _table(file, p.get("sheet"), p.get("header_row", 0))
    cols, sign = p.get("columns", {}), p.get("sign", "positive_expenses")
    c_date = _pick(df, cols, "date", "the date", who, required=True)
    c_cat = _pick(df, cols, "category", "the category", who, required=True)
    if sign == "debit_credit":
        c_deb = _pick(df, cols, "debit", "the debit amount", who, required=True)
        c_cre = _pick(df, cols, "credit", "the credit amount", who, required=True)
    else:
        c_amt = _pick(df, cols, "amount", "the amount", who, required=True)
    c_id, c_desc, c_ven = (_pick(df, cols, f, f, who) for f in ("txn_id", "description", "vendor"))
    rows, skipped = [], []
    for row, r in df.iterrows():
        txt = " ".join(str(r[c]) for c in (c_cat, c_desc, c_ven) if c)
        money = [r[c_deb], r[c_cre]] if sign == "debit_credit" else [r[c_amt]]
        if all(_blank(m) for m in money) and _blank(r[c_date]):
            skipped.append((row, txt.strip()))
            continue
        if _TOTAL_ROW.match(str(r[c_cat]).strip()) or _TOTAL_ROW.match(str(r[c_desc]).strip() if c_desc else ""):
            skipped.append((row, txt.strip()))
            continue
        if sign == "debit_credit":
            cents = (_money(r[c_deb], who, row, "the debit column") or 0) - (_money(r[c_cre], who, row, "the credit column") or 0)
        else:
            v = _money(r[c_amt], who, row, "the amount column")
            if v is None:
                raise IngestError(f"{who} row {row}: the amount is blank.")
            cents = -v if sign == "negative_expenses" else v
        rows.append((row, r, cents))
    dates = _parse_dates([r[c_date] for _, r, _ in rows], p.get("date_order"), who)
    for (row, r, _), d in zip(rows, dates):
        if d is None:
            raise IngestError(f"{who} row {row}: the date is blank.")
    recs = [{
        "txn_id": str(r[c_id]).strip() if c_id and str(r[c_id]).strip() else f"T{i + 1:04d}",
        "date": d, "amount_cents": cents, "category": str(r[c_cat]).strip(),
        "description": str(r[c_desc]).strip() if c_desc else "",
        "vendor": str(r[c_ven]).strip() if c_ven else "", "source_row": row,
    } for i, ((row, r, cents), d) in enumerate(zip(rows, dates))]
    out = pd.DataFrame(recs, columns=ACTUAL_COLS + ["source_row"])
    out["date"] = pd.to_datetime(out["date"])
    out["amount_cents"] = out["amount_cents"].astype("int64")
    out.attrs["skipped_rows"] = skipped
    return out
